"""Creates sample/leads_sample.xlsx with the documented 'Leads' columns.
Replace the example.com addresses with inboxes YOU control before a real test."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ["first_name", "last_name", "full_name", "company_name", "email", "Observed Gaps/Signal",
           "AI/ web Opportunity (Solution)", "Pitch Angle", "consent_status", "unsubscribe_url",
           "send_status", "notes"]
ROWS = [
    ["Sarah", "Khan", "Sarah Khan", "Northstar Realty", "sarah.khan@example.com",
     "Slow lead response and missed website inquiries", "AI lead-response agent connected to CRM",
     "Start with a small pilot to improve response time", "opted_in", "", "Pending", "Warm lead (EXAMPLE ROW - replace)"],
    ["Omar", "Sheikh", "Omar Sheikh", "Harbor Property Group", "omar.sheikh@example.com",
     "Manual WhatsApp follow-ups after viewings", "WhatsApp automation with CRM sync",
     "Automate follow-ups for one property team first", "legitimate_interest", "", "Pending", ""],
    ["Test", "Bad", "Test Bad", "No Consent Co", "noconsent@example.com", "-", "-", "-", "", "", "Pending", "Should be blocked: no consent"],
    ["Test", "Invalid", "Test Invalid", "Bad Email Co", "not-an-email", "-", "-", "-", "opted_in", "", "Pending", "Should be blocked: invalid email"],
    ["Sarah", "Khan", "Sarah Khan", "Northstar Realty", "sarah.khan@example.com", "-", "-", "-", "opted_in", "", "Pending", "Should be blocked: duplicate"],
]
wb = Workbook()
ws = wb.active
ws.title = "Leads"
ws.append(HEADERS)
for r in ROWS:
    ws.append(r)
for c in ws[1]:
    c.font = Font(name="Arial", bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", start_color="1F3864")
    c.alignment = Alignment(wrap_text=True, vertical="center")
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name="Arial")
for i, h in enumerate(HEADERS, 1):
    ws.column_dimensions[get_column_letter(i)].width = max(14, min(42, len(h) + 6))
ws.freeze_panes = "A2"

info = wb.create_sheet("README")
for line in ["Data must stay on the FIRST sheet, named Leads (one row = one recipient).",
             "Required: email. Recommended: full_name, company_name, Observed Gaps/Signal, AI/ web Opportunity (Solution), Pitch Angle, consent_status.",
             "consent_status must be opted_in or legitimate_interest or the row is blocked.",
             "Rows 4-6 are deliberate test failures (no consent / invalid email / duplicate)."]:
    info.append([line])
info.column_dimensions["A"].width = 120
wb.save("sample/leads_sample.xlsx")
print("wrote sample/leads_sample.xlsx")
