"""Pre-flight check of a leads .xlsx BEFORE uploading it to n8n.
Usage: python scripts/validate_leads.py path/to/leads.xlsx"""
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

EMAIL = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
OK_CONSENT = {"opted_in", "legitimate_interest"}
RECOMMENDED = ["first_name", "full_name", "company_name", "Observed Gaps/Signal",
               "AI/ web Opportunity (Solution)", "Pitch Angle", "consent_status"]


def main(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    if ws.title != "Leads":
        print(f"WARNING: first sheet is '{ws.title}', expected 'Leads'.")
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    lower = [h.lower() for h in headers]
    if "email" not in lower:
        sys.exit("ERROR: required column 'email' is missing.")
    for col in RECOMMENDED:
        if col.lower() not in lower:
            print(f"note: recommended column missing: {col}")

    supp = set()
    sp = Path("data/files/suppression.csv")
    if sp.exists():
        supp = {l.split(",")[0].strip().lower() for l in sp.read_text().splitlines()[1:] if l.strip()}

    ei, ci = lower.index("email"), (lower.index("consent_status") if "consent_status" in lower else None)
    seen, ok, bad = set(), 0, []
    for n, row in enumerate(rows[1:], start=2):
        email = str(row[ei] or "").strip().lower()
        reasons = []
        if not EMAIL.match(email):
            reasons.append("invalid email")
        if ci is None or str(row[ci] or "").strip().lower() not in OK_CONSENT:
            reasons.append("no consent")
        if email in supp:
            reasons.append("suppressed")
        if email in seen:
            reasons.append("duplicate")
        seen.add(email)
        if reasons:
            bad.append((n, email or "(blank)", ", ".join(reasons)))
        else:
            ok += 1
    print(f"{ok} row(s) would pass validation, {len(bad)} would be blocked.")
    for n, e, r in bad:
        print(f"  row {n}: {e} -> {r}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
